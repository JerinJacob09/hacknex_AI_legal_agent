import difflib
import hashlib
import html
import io
import json
import re
import traceback
import urllib.request

import streamlit as st

st.set_page_config(page_title="NYAYA-QUANTUM", page_icon="⚖️", layout="wide")

MISSING_DEPS = {}

try:
    from smolagents import ToolCallingAgent, LiteLLMModel, Tool
except Exception as exc:
    ToolCallingAgent = None
    LiteLLMModel = None
    Tool = None
    MISSING_DEPS["smolagents[litellm]"] = str(exc)

try:
    from langchain_chroma import Chroma
except Exception as exc:
    Chroma = None
    MISSING_DEPS["langchain-chroma / chromadb"] = str(exc)

try:
    from langchain_huggingface import HuggingFaceEmbeddings
except Exception as exc:
    HuggingFaceEmbeddings = None
    MISSING_DEPS["langchain-huggingface / sentence-transformers"] = str(exc)

try:
    from langchain_core.documents import Document
except Exception as exc:
    Document = None
    MISSING_DEPS["langchain-core"] = str(exc)

try:
    from langchain_text_splitters import RecursiveCharacterTextSplitter
except Exception as exc:
    RecursiveCharacterTextSplitter = None
    MISSING_DEPS["langchain-text-splitters"] = str(exc)

try:
    from pypdf import PdfReader
except Exception:
    PdfReader = None

try:
    from docx import Document as DocxDocument
except Exception:
    DocxDocument = None

EMBEDDING_MODEL = "BAAI/bge-large-en-v1.5"
PERSIST_DIR = "./nyaya_chroma_store"
DEFAULT_HOST = "http://localhost:11434"
DEFAULT_MODEL = "qwen2.5:7b"
MAX_DRAFT_CHARS = 8000

REM_STYLE = "background-color: #f8d7da; color: #721c24; text-decoration: line-through; padding: 2px 4px; border-radius: 4px;"
INJ_STYLE = "background-color: #d4edda; color: #155724; font-weight: bold; padding: 2px 4px; border-radius: 4px;"

SEVERITY_RANK = {"critical": 0, "important": 1, "minor": 2}

VERIFIED_CITATIONS = {
    "Kesavananda Bharati v. State of Kerala": "(1973) 4 SCC 225",
    "Maneka Gandhi v. Union of India": "(1978) 1 SCC 248",
    "Waman Rao v. Union of India": "(1981) 2 SCC 362",
    "Olga Tellis v. Bombay Municipal Corporation": "(1985) 3 SCC 545",
    "Booz Allen and Hamilton v. SBI Home Finance": "(2011) 5 SCC 532",
    "Shayara Bano v. Union of India": "(2017) 9 SCC 1",
    "K.S. Puttaswamy v. Union of India": "(2017) 10 SCC 1",
    "Vidya Drolia v. Durga Trading Corporation": "(2021) 2 SCC 1",
}

VERIFIED_LIST_TEXT = "\n".join(f"- {name}: {cite}" for name, cite in VERIFIED_CITATIONS.items())

SYSTEM_BASE = f"""You are part of NYAYA-QUANTUM, a fully local Indian legal drafting and litigation-support system.

MANDATORY STATUTORY EXPERTISE: Constitution of India (Articles 13, 14, 19, 21, 32, 136, 226, 227); Bharatiya Nyaya Sanhita, 2023 (BNS); Bharatiya Nagarik Suraksha Sanhita, 2023 (BNSS); Bharatiya Sakshya Adhiniyam, 2023 (BSA); Code of Civil Procedure, 1908 (CPC); Indian Contract Act, 1872; Arbitration and Conciliation Act, 1996; Limitation Act, 1963; Indian Stamp Act, 1899.

LITIGATION PATTERNS: Writ Petitions under Articles 32 and 226, Special Leave Petitions under Article 136, maintainability, territorial jurisdiction, cause of action, alternative remedy, delay and laches, limitation, stamping and admissibility, and binding Supreme Court precedent.

HARD RULES:
1. Never invent facts, dates, amounts, party names, section numbers or case citations.
2. If a critical fact is missing (limitation dates, suit valuation, stamping details, place of cause of action, arbitration seat or venue, impugned order date), write "RFAE:" followed by the evidence needed. Do not assume the fact.
3. Cite a judgment only if it appears in the supplied case-file context or in the verified list below. Otherwise state the principle without a citation and write "citation to be verified".
4. For offences and proceedings after 1 July 2024 use BNS, BNSS and BSA. Flag any reference to the Indian Penal Code, Code of Criminal Procedure or Indian Evidence Act and state the corresponding new statute only where you are certain.
5. Output is a drafting aid for review by a qualified advocate and is not legal advice.

VERIFIED CITATIONS (the only real-world citations you may use without case-file support):
{VERIFIED_LIST_TEXT}
"""

ANALYST_SYSTEM = SYSTEM_BASE + """
ROLE: Document Analyst Agent. You perform precision retrieval, fact extraction and historical case matching strictly from the supplied draft and the retrieved case-file context. Quote or paraphrase only what is present. State clearly when the case files contain no matching authority. Case files labelled FICTIONAL are training material and must never be presented as real precedent."""

SCHOLAR_SYSTEM = SYSTEM_BASE + """
ROLE: Theoretical Legal Scholar Agent. You reason about statutory interpretation, constitutional doctrine, jurisprudential principles and boilerplate clause drafting under Indian law. You work from legal knowledge rather than from case files, so you must be especially strict about citations: give a citation only if it is in the verified list, otherwise name the principle and write "citation to be verified". Separate settled law from contested law."""

CRITIC_SYSTEM = SYSTEM_BASE + """
ROLE: Orchestration Supervisor in strict critique mode. You audit a marked-up legal redraft against the original. Check for: invented facts, invented or unverifiable citations, wrong or obsolete statutes (IPC, CrPC, Indian Evidence Act used where BNS, BNSS, BSA apply), missing RFAE flags for absent critical facts, clauses that are unenforceable or contrary to the Indian Contract Act, Limitation Act, Arbitration and Conciliation Act or Indian Stamp Act, and broken change markers. Respond with ONLY a JSON object of the form {"approved": true or false, "issues": ["issue 1", "issue 2"]}."""

CSS = """
.stApp { background-color: #0F172A; color: #E2E8F0; }
header[data-testid="stHeader"] { background-color: #0F172A; }
section[data-testid="stSidebar"] { background-color: #1E293B; border-right: 1px solid #334155; }
section[data-testid="stSidebar"] * { color: #E2E8F0; }
h1, h2, h3, h4 { color: #38BDF8 !important; }
.stTabs [data-baseweb="tab-list"] { gap: 6px; border-bottom: 1px solid #334155; }
.stTabs [data-baseweb="tab"] { background-color: #1E293B; color: #E2E8F0; border-radius: 8px 8px 0 0; padding: 10px 16px; }
.stTabs [aria-selected="true"] { background-color: #0B3A53; color: #38BDF8; border-bottom: 2px solid #38BDF8; }
.stButton > button { background-color: #1E293B; color: #38BDF8; border: 1px solid #38BDF8; border-radius: 8px; font-weight: 600; }
.stButton > button:hover { background-color: #38BDF8; color: #0F172A; }
.stButton > button[kind="primary"] { background-color: #38BDF8; color: #0F172A; }
.stTextArea textarea, .stTextInput input { background-color: #1E293B; color: #E2E8F0; border: 1px solid #334155; }
.nq-panel { background-color: #1E293B; border: 1px solid #334155; border-radius: 10px; padding: 16px; line-height: 1.75; font-size: 0.95rem; min-height: 200px; max-height: 620px; overflow-y: auto; }
.nq-card { background-color: #1E293B; border-left: 4px solid #38BDF8; border-radius: 8px; padding: 14px 18px; margin-bottom: 14px; }
.nq-card.shield { border-left-color: #FBBF24; }
.nq-card.procedural { border-left-color: #F87171; }
.nq-banner { background-color: #0B3A53; border: 1px solid #38BDF8; border-radius: 8px; padding: 10px 14px; margin-bottom: 12px; font-size: 0.9rem; }
.nq-footer { color: #94A3B8; font-size: 0.8rem; text-align: center; margin-top: 30px; }
"""

SEED_CASES = [
    {
        "id": "seed-mock-001",
        "title": "Meridian Infra Ltd. v. Corvus Logistics Pvt. Ltd. (FICTIONAL TRAINING CASE)",
        "citation": "MOCK/2024/001",
        "text": (
            "FICTIONAL TRAINING CASE, NOT REAL PRECEDENT. Citation: MOCK/2024/001. Subject: indemnification liability dispute. "
            "Meridian Infra engaged Corvus Logistics for cargo handling under a services agreement containing an uncapped indemnity clause "
            "against 'all losses of any kind'. A warehouse fire destroyed goods owned by a third party and Meridian claimed the full "
            "value from Corvus. Held (fictional): an indemnity under Sections 124 and 125 of the Indian Contract Act, 1872 covers loss caused by "
            "the conduct of the promisor or any other person, but the indemnity-holder must prove actual loss and mitigation; an "
            "uncapped clause is enforceable only to the extent of proven loss, and Section 74 compensation principles apply to any "
            "liquidated sum. The claim was time-barred in part because suit was filed beyond three years from the date the loss crystallised "
            "under the Limitation Act, 1963. The agreement was inadmissible until stamp duty and penalty under the Indian Stamp Act, 1899 were paid."
        ),
    },
    {
        "id": "seed-mock-002",
        "title": "Arbor Ledger DAO v. Helix Settlement Services (FICTIONAL TRAINING CASE)",
        "citation": "MOCK/2024/002",
        "text": (
            "FICTIONAL TRAINING CASE, NOT REAL PRECEDENT. Citation: MOCK/2024/002. Subject: smart contract mediation seating dispute. "
            "A smart contract executed on a public ledger referred disputes to 'mediation at the node jurisdiction of the majority validators'. "
            "The parties disagreed on the seat. Held (fictional): under the Arbitration and Conciliation Act, 1996 the seat must be designated "
            "or clearly determinable from the agreement; a floating on-chain reference to validator location is not a designated seat; "
            "the court applied the closest-connection test, treated the place of performance of the off-chain settlement as the seat, "
            "and held that the courts of that place alone had supervisory jurisdiction. The Court recommended that draftsmen of automated contracts "
            "name a fixed seat and venue in the code comments and in the off-chain master agreement, and confirm electronic signatures and stamping."
        ),
    },
    {
        "id": "seed-mock-003",
        "title": "Ravindra Textiles v. State Industrial Corporation (FICTIONAL TRAINING CASE)",
        "citation": "MOCK/2024/003",
        "text": (
            "FICTIONAL TRAINING CASE, NOT REAL PRECEDENT. Citation: MOCK/2024/003. Subject: maintainability of a writ petition in a contractual allotment dispute. "
            "The petitioner challenged cancellation of an industrial plot allotment under Article 226. Held (fictional): a writ is maintainable "
            "against a State instrumentality where the action is arbitrary and violates Article 14, even in contractual matters, but the High Court may "
            "decline relief where an efficacious alternative remedy exists or where disputed questions of fact require evidence. Delay of more than "
            "three years without explanation attracted laches although the Limitation Act, 1963 does not strictly govern writ jurisdiction. "
            "Territorial jurisdiction lies where any part of the cause of action arises under Article 226(2). An appeal under Article 136 "
            "requires the certified copy of the impugned order and a statement of the date of pronouncement."
        ),
    },
]

SAMPLE_DRAFT = (
    "INDEMNITY AND SERVICES AGREEMENT\n\n"
    "This agreement is made between Alpha Traders and Beta Logistics. Beta Logistics shall indemnify Alpha Traders against all losses of any kind "
    "without any limit whatsoever. Any dispute shall be settled by arbitration in a place to be decided later. If Beta Logistics commits any "
    "offence under the Indian Penal Code, Alpha Traders may file an FIR and the case shall be tried under the Code of Criminal Procedure. "
    "Evidence shall be governed by the Indian Evidence Act. Payment shall be made within reasonable time. Alpha Traders may sue Beta Logistics "
    "at any time in any court in India. This agreement is valid even if it is not stamped."
)

ROUTE_TERMS = {
    "analyst": ["case", "precedent", "document", "draft", "uploaded", "facts", "retrieve", "similar", "judgment", "clause", "agreement", "petition"],
    "scholar": ["explain", "doctrine", "interpret", "article", "statute", "section", "define", "jurisprudence", "constitution", "act", "sanhita"],
}


def check_ollama(host, model):
    try:
        with urllib.request.urlopen(host.rstrip("/") + "/api/tags", timeout=3) as response:
            data = json.loads(response.read().decode("utf-8"))
    except Exception as exc:
        return False, f"Ollama is not reachable at {host}: {exc}"
    names = [entry.get("name", "") for entry in data.get("models", [])]
    if model not in names and f"{model}:latest" not in names:
        return False, f"Ollama is running but model '{model}' is not installed. Run: ollama pull {model}"
    return True, f"Ollama online, model '{model}' available."


def read_upload(uploaded):
    name = uploaded.name.lower()
    raw = uploaded.getvalue()
    if not raw:
        return "", f"{uploaded.name} is empty."
    try:
        if name.endswith((".txt", ".md")):
            text = raw.decode("utf-8", errors="replace")
        elif name.endswith(".pdf"):
            if PdfReader is None:
                return "", "PDF support needs the 'pypdf' package (pip install pypdf)."
            reader = PdfReader(io.BytesIO(raw))
            text = "\n".join((page.extract_text() or "") for page in reader.pages)
        elif name.endswith(".docx"):
            if DocxDocument is None:
                return "", "DOCX support needs the 'python-docx' package (pip install python-docx)."
            document = DocxDocument(io.BytesIO(raw))
            text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        else:
            return "", f"Unsupported file type for {uploaded.name}."
    except Exception as exc:
        return "", f"Could not read {uploaded.name}: {exc}"
    if not text.strip():
        return "", f"No extractable text found in {uploaded.name}. Scanned PDFs need OCR first."
    return text, None


@st.cache_resource(show_spinner="Loading BAAI/bge-large-en-v1.5 embeddings locally...")
def load_embeddings(model_name):
    if HuggingFaceEmbeddings is None:
        return None, "langchain-huggingface is not installed."
    try:
        import torch

        device = "cuda" if torch.cuda.is_available() else "cpu"
    except Exception:
        device = "cpu"
    try:
        embeddings = HuggingFaceEmbeddings(
            model_name=model_name,
            model_kwargs={"device": device},
            encode_kwargs={"normalize_embeddings": True},
        )
        embeddings.embed_query("warm-up")
        return embeddings, None
    except Exception as exc:
        return None, f"Embedding model failed to load: {exc}"


def seed_database(vector_store):
    ids = [case["id"] for case in SEED_CASES]
    existing = set(vector_store.get(ids=ids).get("ids", []))
    fresh = [case for case in SEED_CASES if case["id"] not in existing]
    if not fresh:
        return 0
    documents = [
        Document(
            page_content=case["text"],
            metadata={"source": case["title"], "citation": case["citation"], "kind": "seed", "fictional": True},
        )
        for case in fresh
    ]
    vector_store.add_documents(documents, ids=[case["id"] for case in fresh])
    return len(fresh)


@st.cache_resource(show_spinner="Initialising local Chroma vector store and seeding mock cases...")
def load_vectorstore(_embeddings, persist_dir):
    if Chroma is None or Document is None:
        return None, "langchain-chroma or langchain-core is not installed."
    if _embeddings is None:
        return None, "Embeddings are unavailable, so the vector store cannot be built."
    try:
        vector_store = Chroma(
            collection_name="nyaya_quantum",
            embedding_function=_embeddings,
            persist_directory=persist_dir,
        )
        seed_database(vector_store)
        return vector_store, None
    except Exception as exc:
        return None, f"Chroma initialisation or seeding failed: {exc}"


def ingest_text(vector_store, text, source):
    if RecursiveCharacterTextSplitter is None:
        raise RuntimeError("langchain-text-splitters is not installed.")
    splitter = RecursiveCharacterTextSplitter(chunk_size=900, chunk_overlap=120)
    chunks = [chunk for chunk in splitter.split_text(text) if chunk.strip()]
    digest = hashlib.sha256(text.encode("utf-8")).hexdigest()[:16]
    ids = [f"upl-{digest}-{index}" for index in range(len(chunks))]
    existing = set(vector_store.get(ids=ids).get("ids", [])) if ids else set()
    pending = [(chunk_id, chunk) for chunk_id, chunk in zip(ids, chunks) if chunk_id not in existing]
    if pending:
        documents = [
            Document(page_content=chunk, metadata={"source": source, "citation": "uploaded file", "kind": "uploaded", "fictional": False})
            for _, chunk in pending
        ]
        vector_store.add_documents(documents, ids=[chunk_id for chunk_id, _ in pending])
    return len(pending), len(chunks)


def collection_size(vector_store):
    try:
        return len(vector_store.get(include=[]).get("ids", []))
    except Exception:
        return 0


def retrieve_context(vector_store, query, k=4):
    if vector_store is None or not query.strip():
        return "", ""
    try:
        documents = vector_store.similarity_search(query[:1500], k=k)
    except Exception:
        return "", ""
    blocks = []
    for document in documents:
        source = document.metadata.get("source", "unknown")
        citation = document.metadata.get("citation", "n/a")
        blocks.append(f"[SOURCE: {source} | CITATION: {citation}]\n{document.page_content}")
    joined = "\n\n".join(blocks)
    return joined, joined


if Tool is not None:

    class CaseFileSearchTool(Tool):
        name = "search_case_files"
        description = "Searches the local vector database of case files and returns the most relevant passages with source labels."
        inputs = {"query": {"type": "string", "description": "The legal issue, clause, party or fact pattern to search for."}}
        output_type = "string"

        def __init__(self, vector_store):
            super().__init__()
            self.vector_store = vector_store

        def forward(self, query: str) -> str:
            context, _ = retrieve_context(self.vector_store, query, 4)
            return context or "No matching case files found."


@st.cache_resource(show_spinner="Starting local agents...")
def build_agents(_vector_store, model_id, host):
    if ToolCallingAgent is None or LiteLLMModel is None:
        return None, "smolagents with LiteLLM support is not installed."
    try:
        model = LiteLLMModel(model_id=f"ollama_chat/{model_id}", api_base=host, api_key="ollama", num_ctx=8192)
        analyst_tools = [CaseFileSearchTool(_vector_store)] if _vector_store is not None else []
        analyst = ToolCallingAgent(tools=analyst_tools, model=model, max_steps=4, verbosity_level=0)
        scholar = ToolCallingAgent(tools=[], model=model, max_steps=3, verbosity_level=0)
        return {"model": model, "analyst": analyst, "scholar": scholar}, None
    except Exception as exc:
        return None, f"Agent construction failed: {exc}"


def run_agent(agent, model, system, task):
    failure = None
    if agent is not None:
        try:
            output = agent.run(f"{system}\n\nTASK:\n{task}")
            text = output if isinstance(output, str) else str(output)
            if text.strip():
                return text.strip(), None
            failure = "Agent returned an empty answer."
        except Exception as exc:
            failure = f"Agent run failed: {exc}"
    try:
        message = model(
            [
                {"role": "system", "content": [{"type": "text", "text": system}]},
                {"role": "user", "content": [{"type": "text", "text": task}]},
            ]
        )
        text = (getattr(message, "content", "") or "").strip()
        if text:
            return text, failure
        return "", (failure or "") + " Direct model call returned nothing."
    except Exception as exc:
        return "", f"{failure or ''} Direct model call failed: {exc}".strip()


def route_task(text):
    lowered = text.lower()
    scores = {worker: sum(lowered.count(term) for term in terms) for worker, terms in ROUTE_TERMS.items()}
    workers = [worker for worker, score in scores.items() if score > 0]
    if not workers:
        workers = ["scholar"]
    return workers, scores


def extract_json(text):
    if not text:
        return None
    match = re.search(r"\{.*\}", text, re.DOTALL)
    if not match:
        return None
    try:
        return json.loads(match.group(0))
    except Exception:
        return None


def strip_fences(text):
    text = re.sub(r"^```[a-zA-Z]*\s*", "", text.strip())
    return re.sub(r"\s*```$", "", text).strip()


MARK_RE = re.compile(r"\[\[(REM|INJ):(.*?)\]\]", re.DOTALL)
LEGACY_REM_RE = re.compile(r"~~\[REMOVED:\s*(.*?)\]~~", re.DOTALL)
LEGACY_INJ_RE = re.compile(r"\*\*\[INJECTED:\s*(.*?)\]\*\*", re.DOTALL)
CITE_RE = re.compile(
    r"(?:\(\d{4}\)\s*\d+\s*SCC\s*\d+|\[\d{4}\]\s*\d+\s*SCR\s*\d+|\(\d{4}\)\s*\d+\s*SCR\s*\d+|AIR\s*\d{4}\s*SC\s*\d+|\d{4}\s*SCC\s*OnLine\s*[A-Za-z]+\s*\d+|MOCK/\d{4}/\d+)"
)


def normalize_markers(text):
    text = LEGACY_REM_RE.sub(lambda m: f"[[REM:{m.group(1).strip()}]]", text)
    return LEGACY_INJ_RE.sub(lambda m: f"[[INJ:{m.group(1).strip()}]]", text)


def reconstruct_original(marked):
    return MARK_RE.sub(lambda m: m.group(2) if m.group(1) == "REM" else "", marked)


def reconstruct_revised(marked):
    return MARK_RE.sub(lambda m: m.group(2) if m.group(1) == "INJ" else "", marked)


def collapse(text):
    return " ".join(text.split())


def structural_issues(original, marked):
    issues = []
    openers = len(re.findall(r"\[\[(?:REM|INJ):", marked))
    complete = len(MARK_RE.findall(marked))
    if openers != complete:
        issues.append("Some change markers are malformed or nested; every marker must be exactly [[REM:text]] or [[INJ:text]].")
    if complete == 0:
        issues.append("No change markers were produced; mark every deletion and insertion.")
    elif collapse(reconstruct_original(marked)) != collapse(original):
        issues.append("Unmarked changes detected: removing all [[INJ:]] markers and restoring all [[REM:]] text must reproduce the original draft exactly.")
    return issues


def diff_markers(original, revised):
    original_tokens = re.findall(r"\s+|\S+", original)
    revised_tokens = re.findall(r"\s+|\S+", revised)
    matcher = difflib.SequenceMatcher(None, original_tokens, revised_tokens, autojunk=False)
    output = []
    for tag, i1, i2, j1, j2 in matcher.get_opcodes():
        if tag == "equal":
            output.append("".join(original_tokens[i1:i2]))
            continue
        if i2 > i1:
            output.append("[[REM:" + "".join(original_tokens[i1:i2]) + "]]")
        if j2 > j1:
            output.append("[[INJ:" + "".join(revised_tokens[j1:j2]) + "]]")
    return "".join(output)


def norm_cite(citation):
    return re.sub(r"\s+", "", citation).upper()


def verify_citations(text, trusted_text="", suffix=" [citation unverified]"):
    trusted = {norm_cite(cite) for cite in VERIFIED_CITATIONS.values()}
    trusted |= {norm_cite(cite) for cite in CITE_RE.findall(trusted_text)}
    flagged = []

    def replace(match):
        citation = match.group(0)
        if norm_cite(citation) in trusted:
            return citation
        flagged.append(citation)
        return citation + suffix

    return CITE_RE.sub(replace, text), flagged


def render_marked_html(marked):
    safe = html.escape(normalize_markers(marked))

    def replace(match):
        style = REM_STYLE if match.group(1) == "REM" else INJ_STYLE
        return f'<span style="{style}">{match.group(2)}</span>'

    return MARK_RE.sub(replace, safe).replace("\n", "<br>")


def render_plain_html(text):
    return html.escape(text).replace("\n", "<br>")


def rule_based_rfae(draft):
    text = draft.lower()
    items = []

    def add(item, why, severity):
        items.append({"item": item, "why": why, "severity": severity, "source": "rule"})

    if not re.search(r"\b\d{1,2}[./-]\d{1,2}[./-]\d{2,4}\b|\b\d{1,2}(st|nd|rd|th)?\s+(january|february|march|april|may|june|july|august|september|october|november|december)\b|\b(19|20)\d{2}\b", text):
        add("Dates of agreement, breach, cause of action and any notice", "Needed to compute limitation under the Limitation Act, 1963 and to test delay and laches.", "critical")
    if not re.search(r"valuation|value of the suit|suit value|₹|\brs\.?\s*\d|\binr\b|\blakh|\bcrore", text):
        add("Valuation of the suit or subject matter and amounts claimed", "Decides pecuniary jurisdiction, court fee under CPC and stamp duty.", "critical")
    if not re.search(r"stamp", text):
        add("Stamping details: stamp duty paid, denomination, state and date", "Unstamped instruments are inadmissible until duty and penalty are paid under the Indian Stamp Act, 1899.", "critical")
    if not re.search(r"cause of action|jurisdiction|situated|resides|registered office|carries on business|place of (performance|execution)", text):
        add("Place of cause of action, party residences and places of performance", "Needed to fix territorial jurisdiction under CPC Sections 16 to 20 or Article 226(2).", "critical")
    if not re.search(r"petitioner|respondent|plaintiff|defendant|claimant|appellant|party|parties", text):
        add("Full identity, capacity and addresses of all parties", "Needed for cause title, locus standi and service.", "important")
    if "arbitrat" in text and not re.search(r"\bseat\b|\bvenue\b", text):
        add("Arbitration seat and venue, and the full arbitration clause", "Seat decides supervisory courts under the Arbitration and Conciliation Act, 1996.", "critical")
    if "arbitrat" in text and not re.search(r"notice|section 21|appointment", text):
        add("Notice invoking arbitration and appointment mechanism", "Needed to show the arbitration was validly invoked.", "important")
    if re.search(r"article 226|article 32|writ", text) and not re.search(r"alternative remedy|fundamental right|state|instrumentality|public authority|article 12", text):
        add("State action facts and availability of an alternative remedy", "Maintainability of a writ depends on Article 12 status and the efficacious alternative remedy rule.", "critical")
    if re.search(r"article 136|\bslp\b|special leave", text) and not re.search(r"impugned (order|judgment)|certified copy|pronounced", text):
        add("Impugned order or judgment, date of pronouncement and certified copy", "Needed for limitation and for an SLP under Article 136.", "critical")
    if re.search(r"\bfir\b|offence|\bbns\b|accused|police", text) and not re.search(r"fir no|crime no|case no|police station|\bps\b", text):
        add("FIR number, police station, date and sections invoked under BNS", "Needed to test quashing, bail or cognisability under BNSS.", "important")
    if re.search(r"indian penal code|\bipc\b|code of criminal procedure|\bcrpc\b|indian evidence act", text):
        add("Date of the alleged offence or proceeding", "Decides whether BNS, BNSS and BSA (in force from 1 July 2024) or the repealed statutes apply.", "critical")
    return items


def merge_rfae(rule_items, agent_items):
    merged = list(rule_items)
    seen = {re.sub(r"\W+", "", item["item"].lower())[:40] for item in merged}
    for entry in agent_items:
        if not isinstance(entry, dict):
            continue
        item = str(entry.get("item", "")).strip()
        if not item:
            continue
        key = re.sub(r"\W+", "", item.lower())[:40]
        if key in seen:
            continue
        seen.add(key)
        severity = str(entry.get("severity", "important")).lower()
        if severity not in SEVERITY_RANK:
            severity = "important"
        merged.append({"item": item, "why": str(entry.get("why", "")).strip(), "severity": severity, "source": "agent"})
    return sorted(merged, key=lambda entry: SEVERITY_RANK[entry["severity"]])


def parse_strategy(text):
    result = {"primary": "", "alternative": "", "procedural": ""}
    parts = re.split(r"(?im)^\s*#{1,4}\s*(.+?)\s*$", text)
    if len(parts) < 3:
        result["primary"] = text.strip()
        return result
    for index in range(1, len(parts) - 1, 2):
        heading = parts[index].lower()
        body = parts[index + 1].strip()
        if "primary" in heading or "sword" in heading:
            key = "primary"
        elif "alternative" in heading or "shield" in heading:
            key = "alternative"
        elif "procedural" in heading or "maintainab" in heading:
            key = "procedural"
        else:
            continue
        result[key] = (result[key] + "\n\n" + body).strip()
    if not any(result.values()):
        result["primary"] = text.strip()
    return result


def run_pipeline(draft, vector_store, agents, max_rounds, status):
    model = agents["model"]
    analyst = agents["analyst"]
    scholar = agents["scholar"]
    result = {"draft": draft, "log": [], "warnings": []}

    def note(message):
        result["log"].append(message)
        status.write(message)

    workers, scores = route_task(draft)
    if vector_store is None and "analyst" in workers:
        workers.remove("analyst")
        result["warnings"].append("Vector store unavailable: precedent retrieval was skipped.")
    note(f"Supervisor routing: workers = {workers}, keyword scores = {scores}.")

    note("Phase 1: scanning for missing evidence (RFAE).")
    rule_items = rule_based_rfae(draft)
    rfae_task = (
        "List the critical facts missing from this legal draft that a court or opposing counsel would demand. "
        "Respond with ONLY JSON: {\"missing\": [{\"item\": \"...\", \"why\": \"...\", \"severity\": \"critical|important|minor\"}]}.\n\n"
        f"DRAFT:\n{draft}"
    )
    rfae_text, rfae_error = run_agent(scholar, model, SCHOLAR_SYSTEM, rfae_task)
    if rfae_error:
        result["warnings"].append(f"RFAE agent note: {rfae_error}")
    parsed = extract_json(rfae_text) or {}
    agent_items = parsed.get("missing", []) if isinstance(parsed.get("missing", []), list) else []
    result["rfae"] = merge_rfae(rule_items, agent_items)
    note(f"RFAE complete: {len(result['rfae'])} evidence gaps identified.")

    context, context_text = ("", "")
    brief = ""
    if "analyst" in workers:
        note("Phase 2: Document Analyst retrieving and matching case files.")
        context, context_text = retrieve_context(vector_store, draft, 4)
        analyst_task = (
            "Produce a PRECEDENT BRIEF for the draft below. Use only the case-file context. For each relevant case give: source label, "
            "citation exactly as shown in the context, the holding in two sentences, and how it affects the draft. "
            "If nothing is relevant, say so plainly. Remember that FICTIONAL cases are training material only.\n\n"
            f"CASE-FILE CONTEXT:\n{context or 'No case files retrieved.'}\n\nDRAFT:\n{draft}"
        )
        brief, brief_error = run_agent(analyst, model, ANALYST_SYSTEM, analyst_task)
        if brief_error:
            result["warnings"].append(f"Analyst note: {brief_error}")
        brief, _ = verify_citations(brief, context_text)
    result["brief"] = brief

    rfae_lines = "\n".join(f"- {entry['item']}" for entry in result["rfae"]) or "- none identified"

    note("Phase 3: Scholar drafting the mutation.")
    mutation_task = (
        "Transform the raw legal draft below into a legally sound Indian-law draft.\n"
        "OUTPUT FORMAT RULES:\n"
        "- Output the COMPLETE draft as continuous text and nothing else: no commentary, no code fences, no other markdown.\n"
        "- Keep every original word that you do not change exactly as written.\n"
        "- Mark every deletion exactly as [[REM:deleted text]] and every insertion exactly as [[INJ:inserted text]].\n"
        "- A replacement is a [[REM:old]] immediately followed by [[INJ:new]].\n"
        "- Never nest markers and never place the characters ]] inside a marker.\n"
        "- Where a critical fact is missing, insert [[INJ:<RFAE: what is needed>]] instead of inventing the fact.\n"
        "- Cite judgments only from the case-file context or verified list.\n\n"
        f"MISSING EVIDENCE TO FLAG:\n{rfae_lines}\n\nPRECEDENT BRIEF:\n{brief or 'none'}\n\nDRAFT:\n{draft}"
    )
    raw, mutation_error = run_agent(scholar, model, SCHOLAR_SYSTEM, mutation_task)
    if mutation_error:
        result["warnings"].append(f"Mutation note: {mutation_error}")
    marked = normalize_markers(strip_fences(raw)) if raw else ""

    critique_log = []
    approved = False
    for round_number in range(1, max_rounds + 1):
        if not marked:
            critique_log.append(f"Round {round_number}: no draft produced.")
            break
        issues = structural_issues(draft, marked)
        critic_task = (
            f"ORIGINAL DRAFT:\n{draft}\n\nMARKED-UP REDRAFT:\n{marked}\n\n"
            f"CASE-FILE CONTEXT:\n{context_text or 'none'}\n\nAudit the redraft and respond with ONLY the JSON verdict."
        )
        critic_text, critic_error = run_agent(None, model, CRITIC_SYSTEM, critic_task)
        verdict = extract_json(critic_text)
        agent_issues = []
        if critic_error and not critic_text:
            critique_log.append(f"Round {round_number}: critic unavailable ({critic_error}).")
        elif verdict is None:
            critique_log.append(f"Round {round_number}: critic reply was not valid JSON and was ignored.")
        elif not verdict.get("approved", False):
            agent_issues = [str(issue) for issue in verdict.get("issues", [])]
        all_issues = issues + agent_issues
        note(f"Critique round {round_number}/{max_rounds}: {len(all_issues)} issue(s).")
        critique_log.append(f"Round {round_number}: " + ("approved." if not all_issues else "; ".join(all_issues)))
        if not all_issues:
            approved = True
            break
        if round_number == max_rounds:
            break
        rewrite_task = (
            "Rewrite the marked-up draft to resolve every issue listed. Follow the same output format rules exactly: complete draft, "
            "[[REM:...]] and [[INJ:...]] markers only, no nesting, no commentary, no invented facts.\n\n"
            "ISSUES:\n" + "\n".join(f"- {issue}" for issue in all_issues) +
            f"\n\nORIGINAL DRAFT:\n{draft}\n\nPREVIOUS MARKED-UP DRAFT:\n{marked}"
        )
        rewritten, rewrite_error = run_agent(scholar, model, SCHOLAR_SYSTEM, rewrite_task)
        if rewrite_error:
            result["warnings"].append(f"Rewrite note: {rewrite_error}")
        if rewritten:
            marked = normalize_markers(strip_fences(rewritten))

    if not marked:
        marked = draft
        result["warnings"].append("The agents produced no redraft; the original text is shown unchanged.")
    elif structural_issues(draft, marked):
        revised = collapse(reconstruct_revised(marked)) if MARK_RE.search(marked) else marked.strip()
        marked = diff_markers(collapse(draft), revised)
        result["warnings"].append("Agent markers failed the fidelity check, so the change matrix was rebuilt deterministically from the agent's revised text.")
        critique_log.append("Fallback: deterministic word-level diff applied.")
    elif not approved:
        result["warnings"].append("The critique loop ended with unresolved supervisor issues; review the critique log.")

    marked, flagged = verify_citations(marked, context_text, suffix=" <citation unverified>")
    if flagged:
        result["warnings"].append("Unverified citations flagged in the redraft: " + ", ".join(sorted(set(flagged))))
    result["marked"] = marked
    result["revised"] = collapse(reconstruct_revised(marked)) if MARK_RE.search(marked) else marked
    result["critique_log"] = critique_log
    result["approved"] = approved

    note("Phase 4: building the litigation win-strategy.")
    strategy_task = (
        "Prepare an aggressive but honest litigation map for the party who benefits from this draft. Use EXACTLY these three markdown headings:\n"
        "## PRIMARY STATUTORY DIRECTIVES\n## ALTERNATIVE SUBMISSIONS\n## PROCEDURAL MAINTAINABILITY OBJECTIONS\n"
        "Under PRIMARY give the core sword arguments with statutory sections. Under ALTERNATIVE give shield arguments that apply if the primary arguments fail. "
        "Under PROCEDURAL give technical objections to defeat opposing counsel (limitation, territorial jurisdiction, stamping and admissibility, alternative remedy, "
        "maintainability under Articles 32, 226 or 136, arbitration seat). Where a fact is missing write RFAE: and do not assume it. "
        "Cite judgments only from the case-file context or verified list.\n\n"
        f"MISSING EVIDENCE:\n{rfae_lines}\n\nPRECEDENT BRIEF:\n{brief or 'none'}\n\nCASE-FILE CONTEXT:\n{context or 'none'}\n\nREVISED DRAFT:\n{result['revised']}"
    )
    strategy_text, strategy_error = run_agent(scholar, model, SCHOLAR_SYSTEM, strategy_task)
    if strategy_error:
        result["warnings"].append(f"Strategy note: {strategy_error}")
    strategy_text, strategy_flagged = verify_citations(strategy_text, context_text)
    if strategy_flagged:
        result["warnings"].append("Unverified citations flagged in the strategy: " + ", ".join(sorted(set(strategy_flagged))))
    result["strategy"] = parse_strategy(strategy_text) if strategy_text else {"primary": "", "alternative": "", "procedural": ""}
    note("Pipeline complete.")
    return result


def load_sample():
    st.session_state["draft_text"] = SAMPLE_DRAFT


def render_card(title, body, css_class=""):
    st.markdown(f'<div class="nq-card {css_class}"><h4>{html.escape(title)}</h4></div>', unsafe_allow_html=True)
    st.markdown(body if body.strip() else "_No content was generated for this section. Check the agent warnings in Tab 1._")


st.markdown(f"<style>{CSS}</style>", unsafe_allow_html=True)

st.session_state.setdefault("draft_text", "")
st.session_state.setdefault("draft_marker", "")
st.session_state.setdefault("result", None)
st.session_state.setdefault("scholar_chat", [])

with st.sidebar:
    st.markdown("## ⚖️ NYAYA-QUANTUM")
    st.caption("100% local Indian legal multi-agent workspace")
    ollama_host = st.text_input("Ollama host", value=DEFAULT_HOST)
    ollama_model = st.text_input("Ollama model", value=DEFAULT_MODEL)
    critique_rounds = st.slider("Max critique rounds", min_value=1, max_value=3, value=2)

    if MISSING_DEPS:
        st.error("Missing dependencies detected.")
        for package, detail in MISSING_DEPS.items():
            with st.expander(package):
                st.code(detail)
        st.code(
            "pip install streamlit smolagents[litellm] langchain-chroma chromadb langchain-huggingface "
            "sentence-transformers langchain-core langchain-text-splitters pypdf python-docx",
            language="bash",
        )

    ollama_ok, ollama_message = check_ollama(ollama_host, ollama_model)
    if ollama_ok:
        st.success(ollama_message)
    else:
        st.warning(ollama_message)

    embeddings, embedding_error = load_embeddings(EMBEDDING_MODEL)
    if embedding_error:
        st.warning(f"Embedding pipeline unavailable. Retrieval features are disabled, but the Scholar sandbox still works.\n\n{embedding_error}")
    else:
        st.success(f"Embeddings ready: {EMBEDDING_MODEL}")

    vector_store, store_error = load_vectorstore(embeddings, PERSIST_DIR)
    if store_error:
        st.warning(f"Vector store unavailable: {store_error}")
    else:
        st.success(f"Chroma ready with {collection_size(vector_store)} stored chunks.")

    if embedding_error or store_error:
        if st.button("Retry initialisation"):
            st.cache_resource.clear()
            st.rerun()

    st.markdown("---")
    st.markdown("#### Index litigation files")
    kb_files = st.file_uploader("Add reference files to the local knowledge base", type=["txt", "md", "pdf", "docx"], accept_multiple_files=True, key="kb_upload")
    if st.button("Embed into Chroma"):
        if vector_store is None:
            st.error("The vector store is unavailable, so files cannot be embedded.")
        elif not kb_files:
            st.warning("Choose at least one file first.")
        else:
            for kb_file in kb_files:
                text, read_error = read_upload(kb_file)
                if read_error:
                    st.warning(read_error)
                    continue
                try:
                    added, total = ingest_text(vector_store, text, kb_file.name)
                    st.success(f"{kb_file.name}: {added} new chunk(s) added out of {total}.")
                except Exception as exc:
                    st.error(f"Embedding failed for {kb_file.name}: {exc}")

    st.markdown("---")
    st.caption("Seeded mock cases are FICTIONAL training material and not real precedent. Outputs are drafting aids for review by a qualified advocate and are not legal advice.")

agents, agents_error = build_agents(vector_store, ollama_model, ollama_host)

st.markdown("# ⚖️ NYAYA-QUANTUM")
st.caption("Supervisor-Worker legal agents on local Ollama, Chroma and BGE embeddings")

tab1, tab2, tab3, tab4 = st.tabs(
    ["🧬 Document Comparison Mutation Room", "⚔️ Litigation Win-Strategy Room", "🎓 Abstract Scholar Sandbox", "📋 RFAE Checklist (Missing Logs)"]
)

with tab1:
    st.markdown("### Document Comparison Mutation Room")
    uploaded_draft = st.file_uploader("Upload a flawed or raw legal draft (.txt, .md, .pdf, .docx)", type=["txt", "md", "pdf", "docx"], key="draft_upload")
    if uploaded_draft is not None:
        marker = f"{uploaded_draft.name}-{uploaded_draft.size}"
        if st.session_state["draft_marker"] != marker:
            draft_text, draft_error = read_upload(uploaded_draft)
            if draft_error:
                st.warning(draft_error)
            else:
                st.session_state["draft_text"] = draft_text
                st.session_state["draft_marker"] = marker
    st.button("Load sample flawed draft", on_click=load_sample)
    st.text_area("Raw legal draft", key="draft_text", height=220, placeholder="Paste a raw or flawed legal draft here, or upload a file above.")
    execute = st.button("Execute MultiAgent Legal Transformation", type="primary")

    if execute:
        draft = st.session_state["draft_text"].strip()
        if not draft:
            st.warning("The draft is empty. Paste text or upload a file before running the transformation.")
        elif agents is None:
            st.error(f"The agent engine is unavailable: {agents_error}")
        elif not ollama_ok:
            st.error(ollama_message)
        else:
            if len(draft) > MAX_DRAFT_CHARS:
                st.warning(f"The draft is {len(draft)} characters. It was truncated to {MAX_DRAFT_CHARS} so the local 7B model context is not exceeded; split long documents and run them separately.")
                draft = draft[:MAX_DRAFT_CHARS]
            try:
                with st.status("Running the multi-agent pipeline...", expanded=True) as status:
                    st.session_state["result"] = run_pipeline(draft, vector_store, agents, critique_rounds, status)
                    status.update(label="Legal transformation complete", state="complete")
            except Exception as exc:
                st.session_state["result"] = None
                st.error(f"The pipeline failed: {exc}")
                with st.expander("Technical details"):
                    st.code(traceback.format_exc())

    result = st.session_state["result"]
    if result:
        for warning in result["warnings"]:
            st.warning(warning)
        removed = len([m for m in MARK_RE.findall(result["marked"]) if m[0] == "REM"])
        injected = len([m for m in MARK_RE.findall(result["marked"]) if m[0] == "INJ"])
        metric_a, metric_b, metric_c = st.columns(3)
        metric_a.metric("Removals", removed)
        metric_b.metric("Injections", injected)
        metric_c.metric("Supervisor verdict", "Approved" if result["approved"] else "Needs human review")
        left, right = st.columns(2)
        with left:
            st.markdown("#### Original raw text")
            st.markdown(f'<div class="nq-panel">{render_plain_html(result["draft"])}</div>', unsafe_allow_html=True)
        with right:
            st.markdown("#### Mutation matrix")
            st.markdown(f'<div class="nq-panel">{render_marked_html(result["marked"])}</div>', unsafe_allow_html=True)
        download_a, download_b = st.columns(2)
        download_a.download_button("Download clean revised draft", data=result["revised"], file_name="nyaya_revised_draft.txt", mime="text/plain")
        download_b.download_button("Download marked-up draft", data=result["marked"], file_name="nyaya_marked_draft.txt", mime="text/plain")
        with st.expander("Supervisor critique log"):
            for line in result["critique_log"]:
                st.write(line)
        with st.expander("Document Analyst precedent brief"):
            st.markdown(result["brief"] or "_No precedent brief was generated._")

with tab2:
    st.markdown("### Litigation Win-Strategy Room")
    result = st.session_state["result"]
    if not result:
        st.info("Run the transformation in the Mutation Room first. The strategy is built from the revised draft, the precedent brief and the RFAE gaps.")
    else:
        st.markdown('<div class="nq-banner">Strategy is a drafting aid. Verify every statutory reference and citation against official sources before filing.</div>', unsafe_allow_html=True)
        strategy = result["strategy"]
        render_card("🗡️ Primary Statutory Directives (core sword arguments)", strategy["primary"])
        render_card("🛡️ Alternative Submissions (shield arguments if primary fail)", strategy["alternative"], "shield")
        render_card("🚧 Procedural Maintainability Objections (limitation, jurisdiction, stamping)", strategy["procedural"], "procedural")

with tab3:
    st.markdown("### Abstract Scholar Sandbox")
    st.caption("Detached chat with the Theoretical Legal Scholar. It never reads the active document workspace or the vector store.")
    if st.button("Clear conversation"):
        st.session_state["scholar_chat"] = []
        st.rerun()
    history_area = st.container()
    question = st.chat_input("Ask the Scholar, e.g. Explain the doctrine of severability under Article 13 with landmark cases", key="scholar_input")
    if question:
        st.session_state["scholar_chat"].append({"role": "user", "content": question})
        if agents is None:
            st.session_state["scholar_chat"].append({"role": "assistant", "content": f"The agent engine is unavailable: {agents_error}"})
        elif not ollama_ok:
            st.session_state["scholar_chat"].append({"role": "assistant", "content": ollama_message})
        else:
            recent = st.session_state["scholar_chat"][-7:-1]
            transcript = "\n".join(f"{turn['role'].upper()}: {turn['content']}" for turn in recent)
            scholar_task = (
                f"PRIOR CONVERSATION:\n{transcript or 'none'}\n\nQUESTION:\n{question}\n\n"
                "Answer as a rigorous Indian-law scholar. Structure the answer with the constitutional or statutory text, doctrine, landmark cases "
                "(only with verified citations, otherwise write 'citation to be verified'), and the contested points."
            )
            with st.spinner("The Scholar is reasoning..."):
                answer, answer_error = run_agent(agents["scholar"], agents["model"], SCHOLAR_SYSTEM, scholar_task)
            if not answer:
                answer = f"The Scholar could not answer: {answer_error}"
            answer, flagged_cites = verify_citations(answer)
            if flagged_cites:
                answer += "\n\n_Note: citations marked unverified are not in the verified list and must be checked before use._"
            st.session_state["scholar_chat"].append({"role": "assistant", "content": answer})
    with history_area:
        if not st.session_state["scholar_chat"]:
            st.info("No messages yet. Ask a question about Indian statutory or constitutional law.")
        for turn in st.session_state["scholar_chat"]:
            with st.chat_message(turn["role"]):
                st.markdown(turn["content"])

with tab4:
    st.markdown("### RFAE Checklist (Request for Additional Evidence)")
    result = st.session_state["result"]
    if not result:
        st.info("Run the transformation in the Mutation Room to generate the checklist of missing factual evidence.")
    elif not result["rfae"]:
        st.success("No missing critical evidence was detected. A human reviewer should still confirm.")
    else:
        icons = {"critical": "🔴", "important": "🟠", "minor": "🟡"}
        checked = 0
        for index, entry in enumerate(result["rfae"]):
            label = f"{icons[entry['severity']]} {entry['severity'].upper()}: {entry['item']}"
            if st.checkbox(label, key=f"rfae_item_{index}"):
                checked += 1
            if entry["why"]:
                st.caption(entry["why"])
        st.progress(checked / len(result["rfae"]), text=f"{checked} of {len(result['rfae'])} items collected")
        export = "\n".join(f"[ ] {entry['severity'].upper()}: {entry['item']} - {entry['why']}" for entry in result["rfae"])
        st.download_button("Download checklist", data=export, file_name="nyaya_rfae_checklist.txt", mime="text/plain")

st.markdown('<div class="nq-footer">NYAYA-QUANTUM is a local drafting aid. It is not legal advice. Seeded cases are fictional. Verify all law and citations before filing.</div>', unsafe_allow_html=True)