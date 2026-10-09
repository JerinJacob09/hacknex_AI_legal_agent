"""Page-aware chunking, BM25 + dense hybrid retrieval (reciprocal rank fusion), optional cross-encoder rerank.

Works with any LangChain-style vector store that offers similarity_search(query, k, filter=...),
add_documents(docs, ids=...) and get(ids=/where=, include=[...]) - i.e. langchain-chroma. No extra dependencies."""
import io
import math
import re
from collections import Counter

from grounding import Chunk, make_chunk_id
from textutil import normalize_space, tokens

_BM25_CACHE = {}
FETCH_LIMIT = 50000


# ---------------------------------------------------------------- reading files with page numbers
def _paged_blocks(paragraphs, block_chars=3000):
    blocks, current, start = [], [], 1
    for index, para in enumerate(paragraphs, 1):
        current.append(para)
        if sum(len(p) for p in current) >= block_chars:
            blocks.append((None, f"¶{start}-{index}", "\n".join(current)))
            current, start = [], index + 1
    if current:
        blocks.append((None, f"¶{start}-{start + len(current) - 1}", "\n".join(current)))
    return blocks


def read_pages(name, raw):
    """Return (pages, error). pages = [(page_number or None, loc_label, text)]. PDFs keep real page numbers;
    txt/md/docx get paragraph-range labels (they have no fixed pagination)."""
    lowered = name.lower()
    try:
        if lowered.endswith(".pdf"):
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(raw))
            pages = [(i, f"p.{i}", page.extract_text() or "") for i, page in enumerate(reader.pages, 1)]
            pages = [p for p in pages if p[2].strip()]
            if not pages:
                return [], f"No extractable text in {name}. Scanned PDFs need OCR first."
            return pages, None
        if lowered.endswith(".docx"):
            from docx import Document as DocxDocument

            paragraphs = [p.text for p in DocxDocument(io.BytesIO(raw)).paragraphs if p.text.strip()]
        elif lowered.endswith((".txt", ".md")):
            paragraphs = [p for p in re.split(r"\n\s*\n", raw.decode("utf-8", errors="replace")) if p.strip()]
        else:
            return [], f"Unsupported file type for {name}."
    except Exception as exc:
        return [], f"Could not read {name}: {exc}"
    if not paragraphs:
        return [], f"No extractable text found in {name}."
    return _paged_blocks(paragraphs), None


# ---------------------------------------------------------------- chunking
def split_text(text, size=900, overlap=120):
    """Paragraph/sentence-aware splitter. Keeps line breaks (so headings stay separate from the next sentence) and
    never cuts mid-sentence unless a single sentence exceeds `size`."""
    pieces = []  # (separator_before, text)
    for line in re.split(r"\n", text):
        line = line.strip()
        if not line:
            continue
        for n, sentence in enumerate(s for s in re.split(r"(?<=[.!?])\s+", line) if s.strip()):
            pieces.append(("\n" if n == 0 else " ", sentence))
    chunks, current = [], ""
    for sep, piece in pieces:
        while len(piece) > size:
            if current:
                chunks.append(current)
                current = ""
            chunks.append(piece[:size])
            piece = piece[max(size - overlap, 1):]
        if current and len(current) + len(piece) + 1 > size:
            chunks.append(current)
            tail = current[-overlap:] if overlap else ""
            tail = tail[tail.find(" ") + 1:] if " " in tail else ""
            current = (tail + " " + piece).strip() if tail else piece
        else:
            current = (current + sep + piece) if current else piece
    if current:
        chunks.append(current)
    return chunks


def make_chunks(pages, doc, case_id, kind="case", citation="", size=900, overlap=120):
    chunks, index = [], 0
    for page_no, loc, text in pages:
        for piece in split_text(text, size, overlap):
            chunks.append(Chunk(
                id=make_chunk_id(case_id, doc, page_no, index, piece), text=piece, doc=doc, page=page_no,
                loc=loc, case_id=case_id, kind=kind, citation=citation or "n/a",
            ))
            index += 1
    return chunks


def chunk_metadata(chunk):
    # Chroma metadata must be str/int/float/bool: page 0 means "no page number".
    return {"chunk_id": chunk.id, "source": chunk.doc, "doc": chunk.doc, "page": int(chunk.page or 0), "loc": chunk.loc,
            "case_id": chunk.case_id, "kind": chunk.kind, "citation": chunk.citation or "n/a", "fictional": False}


def chunk_from_document(document, fallback_id=""):
    meta = getattr(document, "metadata", {}) or {}
    page = int(meta.get("page") or 0) or None
    return Chunk(
        id=meta.get("chunk_id") or getattr(document, "id", None) or fallback_id or make_chunk_id("", meta.get("source", ""), page, 0, document.page_content),
        text=document.page_content, doc=meta.get("doc") or meta.get("source", ""), page=page, loc=meta.get("loc", ""),
        case_id=meta.get("case_id", ""), kind=meta.get("kind", ""), citation=meta.get("citation", ""),
    )


# ---------------------------------------------------------------- indexing
def invalidate_cache():
    _BM25_CACHE.clear()


def index_chunks(store, chunks, batch=64):
    """Add chunks to the store (skipping ids already present). Returns (new_count, total_count)."""
    try:
        from langchain_core.documents import Document
    except Exception:  # lets the module be tested without langchain installed
        Document = _Doc2

    if not chunks:
        return 0, 0
    ids = [c.id for c in chunks]
    existing = set(store.get(ids=ids).get("ids", []))
    pending = [c for c in chunks if c.id not in existing]
    for start in range(0, len(pending), batch):
        part = pending[start:start + batch]
        store.add_documents([Document(page_content=c.text, metadata=chunk_metadata(c)) for c in part], ids=[c.id for c in part])
    invalidate_cache()
    return len(pending), len(chunks)


def _where(case_ids):
    if not case_ids:
        return None
    case_ids = list(case_ids)
    return {"case_id": case_ids[0]} if len(case_ids) == 1 else {"case_id": {"$in": case_ids}}


def fetch_all(store, case_ids=None, limit=FETCH_LIMIT):
    kwargs = {"include": ["documents", "metadatas"], "limit": limit}
    where = _where(case_ids)
    if where:
        kwargs["where"] = where
    data = store.get(**kwargs)
    out = []
    for cid, text, meta in zip(data.get("ids", []), data.get("documents", []) or [], data.get("metadatas", []) or []):
        if not text:
            continue
        meta = dict(meta or {})
        meta.setdefault("chunk_id", cid)
        out.append(chunk_from_document(_Doc(text, meta), cid))
    return out


class _Doc:
    def __init__(self, text, meta):
        self.page_content, self.metadata = text, meta


class _Doc2:
    def __init__(self, page_content, metadata):
        self.page_content, self.metadata = page_content, metadata


# ---------------------------------------------------------------- BM25
class BM25:
    def __init__(self, chunks, k1=1.5, b=0.75):
        self.chunks = chunks
        self.k1, self.b = k1, b
        self.docs = [Counter(tokens(c.text)) for c in chunks]
        self.lengths = [sum(d.values()) for d in self.docs]
        self.avg = (sum(self.lengths) / len(self.lengths)) if self.lengths else 0.0
        df = Counter()
        for d in self.docs:
            df.update(d.keys())
        n = len(self.docs)
        self.idf = {t: math.log(1 + (n - f + 0.5) / (f + 0.5)) for t, f in df.items()}

    def search(self, query, k=10):
        terms = tokens(query)
        if not terms or not self.docs:
            return []
        scored = []
        for i, doc in enumerate(self.docs):
            score = 0.0
            for t in terms:
                f = doc.get(t)
                if f:
                    denom = f + self.k1 * (1 - self.b + self.b * self.lengths[i] / (self.avg or 1))
                    score += self.idf.get(t, 0.0) * f * (self.k1 + 1) / denom
            if score > 0:
                scored.append((score, i))
        scored.sort(reverse=True)
        out = []
        for score, i in scored[:k]:
            chunk = self.chunks[i]
            out.append(Chunk(**{**chunk.__dict__, "score": score}))
        return out


def _bm25_for(store, case_ids):
    key = (id(store), tuple(case_ids or ()))
    if key not in _BM25_CACHE:
        _BM25_CACHE[key] = BM25(fetch_all(store, case_ids))
    return _BM25_CACHE[key]


# ---------------------------------------------------------------- search
def dense_search(store, query, k=6, case_ids=None):
    kwargs = {}
    where = _where(case_ids)
    if where:
        kwargs["filter"] = where
    try:
        docs = store.similarity_search(query[:1500], k=k, **kwargs)
    except Exception:
        return []
    return [chunk_from_document(d) for d in docs]


def rrf(rankings, k=60):
    scores = {}
    for ranking in rankings:
        for rank, chunk in enumerate(ranking, 1):
            scores[chunk.id] = scores.get(chunk.id, 0.0) + 1.0 / (k + rank)
    return scores


def hybrid_search(store, query, k=6, fetch_k=20, case_ids=None, reranker=None, use_bm25=True):
    """Dense + BM25 fused with reciprocal rank fusion, then (optionally) cross-encoder reranked. Returns list[Chunk]."""
    if store is None or not (query or "").strip():
        return []
    dense = dense_search(store, query, fetch_k, case_ids)
    sparse = []
    if use_bm25:
        try:
            sparse = _bm25_for(store, case_ids).search(query[:1500], fetch_k)
        except Exception:
            sparse = []
    fused = rrf([dense, sparse])
    by_id = {c.id: c for c in sparse}
    by_id.update({c.id: c for c in dense})
    ranked = sorted(by_id.values(), key=lambda c: fused.get(c.id, 0.0), reverse=True)
    for c in ranked:
        c.score = fused.get(c.id, 0.0)
    pool = ranked[:fetch_k]
    if reranker is not None and len(pool) > 1:
        try:
            scores = reranker.predict([(query[:600], c.text) for c in pool])
            for c, s in zip(pool, scores):
                c.score = float(s)
            pool = sorted(pool, key=lambda c: c.score, reverse=True)
        except Exception:
            pass
    return pool[:k]


# ---------------------------------------------------------------- metrics
def recall_at_k(chunks, gold, k=None):
    """gold: list of {'doc': str, 'page': int?} and/or {'text': str}. Returns fraction of gold items hit in the top-k chunks."""
    if not gold:
        return None
    top = chunks[:k] if k else chunks
    hits = 0
    for item in gold:
        for c in top:
            if "text" in item and normalize_space(item["text"]).lower() in normalize_space(c.text).lower():
                hits += 1
                break
            if "doc" in item and c.doc == item["doc"] and (item.get("page") in (None, c.page)):
                hits += 1
                break
    return hits / len(gold)
