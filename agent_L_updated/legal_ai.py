import hashlib
import os
import re
from pathlib import Path

import streamlit as st

import entities
import indiankanoon
import retrieval
from entities import CITE_RE, provisions  # noqa: F401  (shared with the grounding layer; re-exported for callers)


EMBEDDING_CHOICES = {
    "BGE large (general, strong retrieval)": "BAAI/bge-large-en-v1.5",
    "InLegalBERT (Indian legal text)": "law-ai/InLegalBERT",
    "InCaseLawBERT (Indian case law)": "law-ai/InCaseLawBERT",
    "Legal-BERT (English legal text)": "nlpaueb/legal-bert-base-uncased",
}

RERANKER_CHOICES = {
    "Off": "",
    "BGE reranker base (faster)": "BAAI/bge-reranker-base",
    "BGE reranker large (best quality)": "BAAI/bge-reranker-large",
}

DEFAULT_LEGAL_LLM = "hf.co/mradermacher/Legal-Saul-Multiverse-7b-GGUF:Q4_K_M"

NER_MODEL_NAME = "en_legal_ner_trf"
NER_INSTALL_HINT = (
    "pip install spacy && pip install "
    "https://huggingface.co/opennyaiorg/en_legal_ner_trf/resolve/main/en_legal_ner_trf-any-py3-none-any.whl "
    "(check the OpenNyAI GitHub or Hugging Face page if this URL has changed)"
)

SUPPORTED_SUFFIXES = (".txt", ".md", ".pdf", ".docx")

# Landmark list and fictional seed cases were written from memory, not retrieved from a source, so they are off.
REFERENCE_LIST_ENABLED = False
LIBRARY_CASE_ID = "library"  # bulk-ingested real judgments (folder / Hugging Face / Indian Kanoon)
_ACTIVE = {"reranker": None, "ner": None, "embed_fn": None, "case_ids": None}


def set_reranker(reranker):
    _ACTIVE["reranker"] = reranker


def set_ner(nlp):
    _ACTIVE["ner"] = nlp


def set_case_ids(case_ids):
    """Scope retrieval to these case ids (None = whole store)."""
    _ACTIVE["case_ids"] = list(case_ids) if case_ids else None


def set_embedder(embeddings, model_name=""):
    """Enable the embedding tie-break for RFAE duplicates, but only for sentence-embedding models (BGE).
    Raw BERT encoders such as InLegalBERT give uniformly high cosines, which would merge unrelated gaps."""
    _ACTIVE["embed_fn"] = embeddings.embed_documents if embeddings is not None and "bge" in model_name.lower() else None


def embed_fn():
    return _ACTIVE["embed_fn"]


def store_dir(base_dir, embedding_model):
    slug = re.sub(r"[^a-z0-9]+", "_", embedding_model.lower()).strip("_")
    return f"{base_dir}_{slug}"


@st.cache_resource(show_spinner="Loading reranker locally...")
def load_reranker(model_name):
    if not model_name:
        return None, None
    try:
        from sentence_transformers import CrossEncoder

        try:
            import torch

            device = "cuda" if torch.cuda.is_available() else "cpu"
        except Exception:
            device = "cpu"
        return CrossEncoder(model_name, device=device, max_length=512), None
    except Exception as exc:
        return None, f"Reranker failed to load: {exc}"


def retrieve_chunks(vector_store, query, k=4, fetch_k=20, case_ids=None):
    """Hybrid (BM25 + dense, RRF-fused) retrieval, reranked if a reranker is loaded. Returns list[grounding.Chunk]."""
    ids = case_ids if case_ids is not None else _ACTIVE["case_ids"]
    return retrieval.hybrid_search(vector_store, query, k=k, fetch_k=fetch_k, case_ids=ids, reranker=_ACTIVE["reranker"])


def format_context(chunks):
    blocks = []
    for chunk in chunks:
        where = f"p.{chunk.page}" if chunk.page else (chunk.loc or "n/a")
        blocks.append(f"[SOURCE: {chunk.doc or 'unknown'} | LOCATION: {where} | CHUNK: {chunk.id} | CITATION: {chunk.citation or 'n/a'}]\n{chunk.text}")
    return "\n\n".join(blocks)


def retrieve_reranked(vector_store, query, k=4, fetch_k=20):
    """Backwards-compatible string interface used by the agents and pipeline."""
    chunks = retrieve_chunks(vector_store, query, k=k, fetch_k=fetch_k)
    joined = format_context(chunks)
    return joined, joined


@st.cache_resource(show_spinner="Loading Indian legal NER model...")
def load_ner():
    try:
        import spacy

        return spacy.load(NER_MODEL_NAME), None
    except Exception as exc:
        return None, f"Legal NER unavailable ({exc}). Install with: {NER_INSTALL_HINT}"


def ner_extract(nlp, text):
    if nlp is None or not text.strip():
        return {}
    try:
        doc = nlp(text[:100000])
    except Exception:
        return {}
    grouped = {}
    for ent in doc.ents:
        value = ent.text.strip()
        if value:
            grouped.setdefault(ent.label_, []).append(value)
    return grouped


def ner_extract_active(text):
    return ner_extract(_ACTIVE["ner"], text)


def enhance_rfae(draft, rule_items):
    nlp = _ACTIVE["ner"]
    if nlp is None:
        return rule_items
    entities = ner_extract(nlp, draft)
    lowered = draft.lower()
    items = list(rule_items)

    def add(item, why, severity):
        items.append({"item": item, "why": why, "severity": severity, "source": "ner"})

    litigation = re.search(r"petition|suit|appeal|writ|plaint|complaint|\bslp\b", lowered)
    if litigation and not entities.get("CASE_NUMBER"):
        add("Case or filing number of the proceeding", "No case number was detected in the draft; courts require it in the cause title.", "important")
    if litigation and not entities.get("COURT"):
        add("Court or forum before which the matter is filed", "No court was detected in the draft; needed to test jurisdiction and maintainability.", "important")
    if litigation and not (entities.get("PETITIONER") or entities.get("RESPONDENT")):
        add("Named petitioner(s) and respondent(s)", "No party names were detected in the draft.", "important")
    return items


def new_provisions(original, revised, context=""):
    return sorted(provisions(revised) - provisions(original) - provisions(context))


def read_file(path):
    path = Path(path)
    suffix = path.suffix.lower()
    try:
        if suffix in (".txt", ".md"):
            return path.read_bytes().decode("utf-8", errors="replace")
        if suffix == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(str(path))
            return "\n".join((page.extract_text() or "") for page in reader.pages)
        if suffix == ".docx":
            from docx import Document as DocxDocument

            return "\n".join(p.text for p in DocxDocument(str(path)).paragraphs)
    except Exception:
        return ""
    return ""


def guess_citation(text):
    match = CITE_RE.search(text[:4000])
    return re.sub(r"\s+", " ", match.group(0)) if match else "n/a"


def ingest_pages(vector_store, pages, doc, case_id, kind="judgment", citation=None):
    full_text = "\n".join(text for _, _, text in pages)
    chunks = retrieval.make_chunks(pages, doc, case_id, kind=kind, citation=citation or guess_citation(full_text), size=1200, overlap=150)
    return retrieval.index_chunks(vector_store, chunks)


def ingest_judgment(vector_store, text, source, citation=None):
    return ingest_pages(vector_store, [(None, "¶1", text)], str(source)[:200], LIBRARY_CASE_ID, "judgment", citation)


def ingest_case_files(vector_store, files, case_id, progress=None):
    """files = [(name, raw_bytes)]. Chunks keep document name and page (PDF) or paragraph range (txt/md/docx)."""
    added = total = 0
    problems = []
    for index, (name, raw) in enumerate(files):
        pages, error = retrieval.read_pages(name, raw)
        if error:
            problems.append(error)
        else:
            new, count = retrieval.index_chunks(vector_store, retrieval.make_chunks(pages, name, case_id, kind="case"))
            added += new
            total += count
        if progress:
            progress((index + 1) / len(files), f"{index + 1}/{len(files)} files")
    return added, total, problems


def purge_reference_entries(vector_store):
    """Remove the fictional seed cases and the from-memory landmark index notes that older versions wrote into the store."""
    if vector_store is None or REFERENCE_LIST_ENABLED:
        return 0
    try:
        stale = vector_store.get(where={"kind": {"$in": ["seed", "landmark"]}}, include=[]).get("ids", [])
        if stale:
            vector_store.delete(ids=stale)
            retrieval.invalidate_cache()
        return len(stale)
    except Exception:
        return 0


def ingest_folder(vector_store, folder, max_files, progress=None):
    root = Path(folder).expanduser()
    if not root.is_dir():
        raise ValueError(f"'{folder}' is not a folder.")
    files = [p for p in sorted(root.rglob("*")) if p.suffix.lower() in SUPPORTED_SUFFIXES][:max_files]
    if not files:
        raise ValueError("No .txt, .md, .pdf or .docx files found in that folder.")
    added = total = 0
    skipped = []
    for index, path in enumerate(files):
        pages, error = retrieval.read_pages(path.name, path.read_bytes())
        if error:
            skipped.append(path.name)
        else:
            new, count = ingest_pages(vector_store, pages, path.name, LIBRARY_CASE_ID)
            added += new
            total += count
        if progress:
            progress((index + 1) / len(files), f"{index + 1}/{len(files)} files")
    return added, total, len(files), skipped


def ingest_hf_dataset(vector_store, dataset_name, split, text_field, title_field, max_rows, config=None, progress=None):
    from datasets import load_dataset

    stream = load_dataset(dataset_name, config or None, split=split, streaming=True)
    added = total = rows = 0
    for index, row in enumerate(stream):
        if index >= max_rows:
            break
        if index == 0 and text_field not in row:
            raise ValueError(f"Column '{text_field}' not found. Available columns: {', '.join(row.keys())}")
        text = row.get(text_field)
        if not isinstance(text, str) or not text.strip():
            continue
        title = str(row.get(title_field)) if title_field and row.get(title_field) else f"{dataset_name}#{index}"
        new, count = ingest_judgment(vector_store, text[:300000], title)
        added += new
        total += count
        rows += 1
        if progress:
            progress((index + 1) / max_rows, f"{index + 1}/{max_rows} rows")
    return added, total, rows


LANDMARK_CASES = [
    ("A.K. Gopalan v. State of Madras", "AIR 1950 SC 27", "Read Article 21 narrowly, treating procedure established by law as any validly enacted procedure; later displaced by Maneka Gandhi."),
    ("State of West Bengal v. Anwar Ali Sarkar", "AIR 1952 SC 75", "Early Article 14 case on reasonable classification and equal protection."),
    ("E.P. Royappa v. State of Tamil Nadu", "(1974) 4 SCC 3", "Arbitrariness is antithetical to equality under Articles 14 and 16."),
    ("Indira Gandhi v. Raj Narain", "1975 Supp SCC 1", "Applied the basic structure doctrine; free and fair elections form part of it."),
    ("Mohinder Singh Gill v. Chief Election Commissioner", "(1978) 1 SCC 405", "The validity of a statutory order must be judged by the reasons in the order itself and cannot be supplemented by later affidavits."),
    ("Minerva Mills v. Union of India", "(1980) 3 SCC 625", "Limited amending power is itself part of the basic structure; balance between Fundamental Rights and Directive Principles."),
    ("Hussainara Khatoon v. Home Secretary, State of Bihar", "(1980) 1 SCC 81", "Speedy trial is part of the right to life and personal liberty under Article 21."),
    ("Ajay Hasia v. Khalid Mujib Sehravardi", "(1981) 1 SCC 722", "Tests for deciding whether a body is an instrumentality of the State under Article 12."),
    ("Mohd. Ahmed Khan v. Shah Bano Begum", "(1985) 2 SCC 556", "A divorced Muslim woman can claim maintenance under Section 125 of the Code of Criminal Procedure."),
    ("Kihoto Hollohan v. Zachillhu", "1992 Supp (2) SCC 651", "Upheld the Tenth Schedule; the Speaker's decision on disqualification is subject to judicial review."),
    ("State of Haryana v. Bhajan Lal", "1992 Supp (1) SCC 335", "Categories of cases in which an FIR or criminal proceeding may be quashed under inherent powers."),
    ("S.R. Bommai v. Union of India", "(1994) 3 SCC 1", "Proclamation under Article 356 is judicially reviewable; secularism is part of the basic structure."),
    ("Joginder Kumar v. State of U.P.", "(1994) 4 SCC 260", "Arrest must not be routine; police need justification beyond the mere power to arrest."),
    ("D.K. Basu v. State of West Bengal", "(1997) 1 SCC 416", "Mandatory guidelines to prevent custodial torture and to regulate arrest and detention."),
    ("Vishaka v. State of Rajasthan", "(1997) 6 SCC 241", "Guidelines against sexual harassment at the workplace drawn from Articles 14, 15, 19 and 21."),
    ("Whirlpool Corporation v. Registrar of Trade Marks", "(1998) 8 SCC 1", "A High Court may entertain a writ despite an alternative remedy where fundamental rights or natural justice are violated, the authority acts without jurisdiction, or vires is challenged."),
    ("I.R. Coelho v. State of Tamil Nadu", "(2007) 2 SCC 1", "Laws placed in the Ninth Schedule remain open to review on the basic structure and fundamental rights."),
    ("A.V. Papayya Sastry v. Government of Andhra Pradesh", "(2007) 4 SCC 221", "A judgment or order obtained by fraud is a nullity and can be challenged at any stage."),
    ("Lalita Kumari v. Government of Uttar Pradesh", "(2014) 2 SCC 1", "Registration of an FIR is mandatory when information discloses a cognizable offence; preliminary inquiry only in limited categories."),
    ("Arnesh Kumar v. State of Bihar", "(2014) 8 SCC 273", "Safeguards against automatic arrest for offences punishable up to seven years."),
    ("Anvar P.V. v. P.K. Basheer", "(2014) 10 SCC 473", "Electronic records as secondary evidence need the certificate under Section 65B(4) of the Indian Evidence Act, now Section 63 of the Bharatiya Sakshya Adhiniyam."),
    ("Kailash Nath Associates v. Delhi Development Authority", "(2015) 4 SCC 136", "Section 74 of the Indian Contract Act: reasonable compensation not exceeding the stipulated sum; loss must be proved unless it is difficult to assess."),
    ("Shreya Singhal v. Union of India", "(2015) 5 SCC 1", "Section 66A of the Information Technology Act struck down as violating Article 19(1)(a)."),
    ("Supreme Court Advocates-on-Record Association v. Union of India", "(2016) 5 SCC 1", "Struck down the National Judicial Appointments Commission Act and the 99th Constitutional Amendment."),
    ("Subramanian Swamy v. Union of India", "(2016) 7 SCC 221", "Upheld criminal defamation as a reasonable restriction on free speech."),
    ("Indus Mobile Distribution v. Datawind Innovations", "(2017) 7 SCC 678", "Designation of a seat of arbitration gives the courts at that seat exclusive jurisdiction over arbitration-related applications."),
    ("Joseph Shine v. Union of India", "(2019) 3 SCC 39", "Section 497 of the Indian Penal Code (adultery) struck down."),
    ("Mayavati Trading v. Pradyuat Deb Burman", "(2019) 8 SCC 714", "At the Section 11 stage the court examines only the existence of an arbitration agreement."),
    ("Ssangyong Engineering and Construction v. National Highways Authority of India", "(2019) 15 SCC 131", "Narrow scope of the public policy ground for setting aside an award under Section 34 after the 2015 amendment."),
    ("Indian Young Lawyers Association v. State of Kerala", "(2019) 11 SCC 1", "Sabarimala temple entry case on equality and freedom of religion."),
    ("Perkins Eastman Architects v. HSCC (India)", "(2020) 20 SCC 760", "A person with an interest in the dispute cannot unilaterally appoint the sole arbitrator."),
    ("Arjun Panditrao Khotkar v. Kailash Kushanrao Gorantyal", "(2020) 7 SCC 1", "The certificate under Section 65B(4) of the Indian Evidence Act is a condition for admitting electronic records as secondary evidence."),
    ("Anuradha Bhasin v. Union of India", "(2020) 3 SCC 637", "Restrictions on internet access and movement must satisfy proportionality and be published."),
    ("Satender Kumar Antil v. Central Bureau of Investigation", "(2022) 10 SCC 51", "Guidelines on bail, categorising offences and directing against needless arrest."),
    ("Navtej Singh Johar v. Union of India", "(2018) 10 SCC 1", "Section 377 of the Indian Penal Code struck down insofar as it criminalised consensual same-sex conduct between adults."),
    ("K.S. Puttaswamy v. Union of India (Aadhaar)", "(2019) 1 SCC 1", "Upheld the Aadhaar Act subject to limits on its use by private entities."),
    ("Bharat Aluminium Co. v. Kaiser Aluminium Technical Services", "(2012) 9 SCC 552", "Part I of the Arbitration and Conciliation Act applies only to arbitrations seated in India."),
    ("Fateh Chand v. Balkishan Dass", "AIR 1963 SC 1405", "Section 74 of the Indian Contract Act allows only reasonable compensation not exceeding the stipulated sum, whether it is a penalty or liquidated damages."),
]

_BASE_VERIFIED = {
    "Kesavananda Bharati v. State of Kerala",
    "Maneka Gandhi v. Union of India",
    "Waman Rao v. Union of India",
    "Olga Tellis v. Bombay Municipal Corporation",
    "Booz Allen and Hamilton v. SBI Home Finance",
    "Shayara Bano v. Union of India",
    "K.S. Puttaswamy v. Union of India",
    "Vidya Drolia v. Durga Trading Corporation",
}

LANDMARK_CITATIONS = {name: cite for name, cite, _ in LANDMARK_CASES if name not in _BASE_VERIFIED}


def seed_landmarks(vector_store):
    from langchain_core.documents import Document

    ids = ["landmark-" + re.sub(r"[^a-z0-9]+", "-", name.lower()).strip("-") for name, _, _ in LANDMARK_CASES]
    existing = set(vector_store.get(ids=ids).get("ids", []))
    documents = []
    new_ids = []
    for cid, (name, cite, principle) in zip(ids, LANDMARK_CASES):
        if cid in existing:
            continue
        text = (
            f"REFERENCE INDEX ENTRY. Case: {name}. Citation: {cite}. Principle: {principle} "
            "This is a short index note, not the judgment; verify against the official report before relying on it."
        )
        documents.append(Document(page_content=text, metadata={"source": name, "citation": cite, "kind": "landmark", "fictional": False}))
        new_ids.append(cid)
    if documents:
        vector_store.add_documents(documents, ids=new_ids)
    return len(documents)


def sidebar_model_settings():
    st.markdown("#### Legal AI models")
    embedding_label = st.selectbox("Embedding model", list(EMBEDDING_CHOICES), index=0, key="nq_embed_choice")
    reranker_label = st.selectbox("Reranker", list(RERANKER_CHOICES), index=0, key="nq_rerank_choice")
    legal_llm = st.text_input(
        "Legal LLM for drafting (Ollama name, blank = use main model)",
        value=DEFAULT_LEGAL_LLM,
        key="nq_legal_llm",
        help="Run `ollama pull " + DEFAULT_LEGAL_LLM + "` first. If that quant tag is not available, open the "
        "Hugging Face repo, pick a listed quant and change the tag. This model is trained mostly on US/EU/UK law, "
        "so keep the verified-citation guards on.",
    )
    use_ner = st.checkbox("Use OpenNyAI legal NER", value=False, key="nq_use_ner")
    st.caption("Changing the embedding model switches to a separate Chroma store, so you may need to re-index your files.")
    return EMBEDDING_CHOICES[embedding_label], RERANKER_CHOICES[reranker_label], legal_llm.strip(), use_ner


def sidebar_case_settings():
    st.markdown("#### Active case")
    case_name = st.text_input("Case name (documents you index below belong to this case)", value="case-1", key="nq_case_id")
    case_id = re.sub(r"[^a-z0-9_-]+", "-", case_name.lower()).strip("-") or "case-1"
    include_library = st.checkbox("Also search the real-judgment library for precedents (drafting only)", value=True, key="nq_use_library")
    st.caption("Case Chat and Case Review only ever use documents of the active case.")
    return case_id, include_library


def sidebar_kanoon(vector_store):
    with st.expander("Indian Kanoon API (real judgments)"):
        st.caption("Needs an API token and internet. Check Indian Kanoon's terms and pricing first. Results go into the judgment library, with the real title as the source.")
        token = st.text_input("API token", value=os.environ.get("INDIANKANOON_TOKEN", ""), type="password", key="nq_ik_token")
        query = st.text_input("Search query", key="nq_ik_query", placeholder="anticipatory bail section 438 economic offence")
        limit = st.number_input("Max judgments", min_value=1, max_value=25, value=5, key="nq_ik_max")
        if st.button("Search and index", key="nq_ik_go"):
            if vector_store is None:
                st.error("The vector store is unavailable.")
            elif not token or not query.strip():
                st.warning("Enter a token and a query.")
            else:
                try:
                    client = indiankanoon.IndianKanoon(token)
                    hits = client.search(query.strip())[: int(limit)]
                    added = 0
                    for hit in hits:
                        doc = client.fetch(hit["tid"])
                        if doc["text"]:
                            new, _ = ingest_pages(vector_store, [(None, "¶1", doc["text"][:300000])], doc["title"][:200], LIBRARY_CASE_ID, "judgment", doc["citation"] or None)
                            added += new
                    st.success(f"{len(hits)} judgments fetched, {added} new chunks indexed.")
                except Exception as exc:
                    st.error(f"Indian Kanoon request failed: {exc}")


def sidebar_data_tools(vector_store):
    st.markdown("---")
    st.markdown("#### Real Indian legal data (judgment library)")
    if vector_store is None:
        st.caption("Vector store unavailable, so bulk ingestion is disabled.")
        return
    sidebar_kanoon(vector_store)
    st.caption("Embedding thousands of judgments on CPU is slow. Start with 20 to 50 files to test, then scale up.")

    with st.expander("Ingest a folder of judgments"):
        folder = st.text_input("Folder path", key="nq_folder_path", placeholder="C:/data/sc_judgments")
        max_files = st.number_input("Max files", min_value=1, max_value=5000, value=50, key="nq_folder_max")
        if st.button("Ingest folder", key="nq_folder_go"):
            bar = st.progress(0.0)
            try:
                added, total, n_files, skipped = ingest_folder(vector_store, folder, int(max_files), lambda f, t: bar.progress(f, text=t))
                st.success(f"{n_files} files processed: {added} new chunks added out of {total}.")
                if skipped:
                    st.warning(f"No extractable text in {len(skipped)} file(s): " + ", ".join(skipped[:5]))
            except Exception as exc:
                st.error(f"Folder ingestion failed: {exc}")

    with st.expander("Stream a Hugging Face dataset"):
        st.caption("Needs `pip install datasets`. Check the dataset card for column names and licence before using.")
        dataset_name = st.text_input("Dataset name", key="nq_hf_name", placeholder="owner/dataset-name")
        config = st.text_input("Config (optional)", key="nq_hf_config")
        split = st.text_input("Split", value="train", key="nq_hf_split")
        text_field = st.text_input("Text column", value="text", key="nq_hf_text")
        title_field = st.text_input("Title column (optional)", key="nq_hf_title")
        max_rows = st.number_input("Max rows", min_value=1, max_value=20000, value=50, key="nq_hf_max")
        if st.button("Stream and embed", key="nq_hf_go"):
            if not dataset_name.strip():
                st.warning("Enter a dataset name first.")
            else:
                bar = st.progress(0.0)
                try:
                    added, total, rows = ingest_hf_dataset(
                        vector_store, dataset_name.strip(), split.strip(), text_field.strip(), title_field.strip(),
                        int(max_rows), config.strip() or None, lambda f, t: bar.progress(f, text=t),
                    )
                    st.success(f"{rows} rows ingested: {added} new chunks added out of {total}.")
                except Exception as exc:
                    st.error(f"Dataset ingestion failed: {exc}")
